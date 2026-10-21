"""Write a DataFrame to Excel using the formatting of an existing Excel template.

This replaces ``StyleFrame.read_excel_as_template(...).to_excel(...)`` (StyleFrame is no longer maintained and does
not support pandas 3) using only pandas and openpyxl, and reproduces StyleFrame's output:

- Header and data cells that fall within the template's bounds take the template cell's formatting, converted the way
  StyleFrame converted it: only RGB colors are kept (theme/indexed colors become a white fill and a black font), border
  colors are black, the fill background color and diagonal borders are dropped, and only these font attributes are
  kept: name, size, color, bold, italic, underline and strikethrough.
- Cells outside the template's bounds get StyleFrame's default style (Arial 12, thin black borders, centered, wrapped,
  white solid fill, unlocked).
- Date/time values use StyleFrame's default date/time number formats.
- Column widths and the header/template row heights are copied from the template.
- An auto-filter can be added to a row.
"""

import datetime as dt
import re
from pathlib import Path
from typing import NamedTuple

import pandas as pd
from openpyxl import load_workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter

_HEX_COLOR = re.compile(r"^([A-Fa-f0-9]{8}|[A-Fa-f0-9]{6})$")
_WHITE = "00FFFFFF"
_BLACK = "00000000"
OUTPUT_SHEET_NAME = "Sheet1"

_DATE_FORMAT = "DD/MM/YY"
_TIME_FORMAT = "HH:MM"
_DATE_TIME_FORMAT = "DD/MM/YY HH:MM"
_NUMBER_FORMAT_BY_TYPE = {
    pd.Timestamp: _DATE_TIME_FORMAT,
    dt.datetime: _DATE_TIME_FORMAT,
    dt.date: _DATE_FORMAT,
    dt.time: _TIME_FORMAT,
}


class CellStyle(NamedTuple):
    bg_color: str = _WHITE
    bold: bool = False
    italic: bool = False
    strikethrough: bool = False
    font: str = "Arial"
    font_size: float = 12.0
    font_color: str = _BLACK
    underline: str | None = None
    number_format: str = "General"
    locked: bool = False
    border_top: str | None = "thin"
    border_right: str | None = "thin"
    border_bottom: str | None = "thin"
    border_left: str | None = "thin"
    horizontal_alignment: str | None = "center"
    vertical_alignment: str | None = "center"
    wrap_text: bool = True
    shrink_to_fit: bool = True
    fill_pattern_type: str | None = "solid"
    indent: float = 0.0
    text_rotation: int = 0

    @classmethod
    def from_cell(cls, cell: Cell) -> "CellStyle":
        def rgb_or(color: Color | None, default: str) -> str:
            rgb = getattr(color, "rgb", None)
            # openpyxl returns a (non-hex) message string as `rgb` for theme and indexed colors
            return rgb if isinstance(rgb, str) and _HEX_COLOR.match(rgb) else default

        font, alignment, border = cell.font, cell.alignment, cell.border
        return cls(
            bg_color=rgb_or(cell.fill.fgColor, _WHITE),
            bold=font.bold,
            italic=font.italic,
            strikethrough=font.strikethrough,
            font=font.name,
            font_size=font.size,
            font_color=rgb_or(font.color, _BLACK),
            underline=font.underline,
            number_format=cell.number_format,
            locked=cell.protection.locked,
            border_top=border.top.border_style,
            border_right=border.right.border_style,
            border_bottom=border.bottom.border_style,
            border_left=border.left.border_style,
            horizontal_alignment=alignment.horizontal,
            vertical_alignment=alignment.vertical,
            wrap_text=alignment.wrap_text or False,
            shrink_to_fit=alignment.shrink_to_fit,
            fill_pattern_type=cell.fill.patternType,
            indent=alignment.indent,
            text_rotation=alignment.text_rotation,
        )

    def apply_to(self, cell: Cell, value: object) -> None:
        cell.font = Font(
            name=self.font,
            size=self.font_size,
            color=Color(self.font_color),
            bold=self.bold,
            italic=self.italic,
            strikethrough=self.strikethrough,
            underline=self.underline,
        )
        cell.fill = PatternFill(patternType=self.fill_pattern_type, fgColor=self.bg_color)
        cell.border = Border(
            top=Side(border_style=self.border_top, color=_BLACK),
            right=Side(border_style=self.border_right, color=_BLACK),
            bottom=Side(border_style=self.border_bottom, color=_BLACK),
            left=Side(border_style=self.border_left, color=_BLACK),
        )
        cell.alignment = Alignment(
            horizontal=self.horizontal_alignment,
            vertical=self.vertical_alignment,
            wrap_text=self.wrap_text,
            shrink_to_fit=self.shrink_to_fit,
            indent=self.indent,
            text_rotation=self.text_rotation,
        )
        cell.number_format = _NUMBER_FORMAT_BY_TYPE.get(type(value), self.number_format)
        cell.protection = Protection(locked=self.locked)


def write_excel_from_template(
    template_path: Path,
    template_sheet_name: str,
    df: pd.DataFrame,
    out_file: Path,
    row_to_add_filters: int | None = None,
) -> None:
    """Write `df` (with a header row, without the index) to the `OUTPUT_SHEET_NAME` sheet of `out_file`, formatted like
    the template sheet.

    `df` is expected to be the template sheet as read by `pd.read_excel` (a single header row), with rows appended.

    :param row_to_add_filters: zero-based row (0 is the header row) to add an auto-filter to, if any
    """
    template = load_workbook(template_path)[template_sheet_name]
    template_num_of_rows, template_num_of_cols = pd.read_excel(template_path, sheet_name=template_sheet_name).shape
    default_style = CellStyle()

    with pd.ExcelWriter(out_file, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=OUTPUT_SHEET_NAME, index=False, na_rep="")
        sheet = writer.sheets[OUTPUT_SHEET_NAME]
        sheet.sheet_view.rightToLeft = False

        # header row is row 1, data starts at row 2 in both the template and the output
        for col, column_name in enumerate(df.columns, start=1):
            for row, value in enumerate([column_name, *df.iloc[:, col - 1]], start=1):
                in_template = col <= template_num_of_cols and row <= template_num_of_rows + 1
                style = CellStyle.from_cell(template.cell(row=row, column=col)) if in_template else default_style
                style.apply_to(sheet.cell(row=row, column=col), value)

        for col in range(1, min(len(df.columns), template_num_of_cols) + 1):
            letter = get_column_letter(col)
            sheet.column_dimensions[letter].width = template.column_dimensions[letter].width

        for row in range(1, min(len(df.index), template_num_of_rows) + 2):
            sheet.row_dimensions[row].height = template.row_dimensions[row].height

        if row_to_add_filters is not None:
            filter_row = row_to_add_filters + 1
            sheet.auto_filter.ref = f"A{filter_row}:{get_column_letter(len(df.columns))}{filter_row}"
