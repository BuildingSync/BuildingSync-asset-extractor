import tempfile
import unittest
from pathlib import Path

import pandas as pd
import pytest
from lxml import etree
from openpyxl import load_workbook
from pandas.testing import assert_frame_equal

from buildingsync_asset_extractor.cts import building_sync_to_cts
from buildingsync_asset_extractor.cts.classes import FacilityAppearance
from buildingsync_asset_extractor.cts.cts import aggregate_facilities

BLANK_CTS_FILE_PATH = (
    Path(__file__).parents[2] / "buildingsync_asset_extractor" / "cts" / "CTS Comprehensive Evaluation Upload Template_20240312_021125.xlsx"
)
BLANK_CTS = pd.read_excel(BLANK_CTS_FILE_PATH, sheet_name="Evaluation Upload Template")

PRIMARYSCHOOL_1_FILE_PATH = Path(__file__).parents[1] / "files" / "PrimarySchool-1.xml"
PRIMARYSCHOOL_2_FILE_PATH = Path(__file__).parents[1] / "files" / "PrimarySchool-2.xml"
OFFICE_3_FILE_PATH = Path(__file__).parents[1] / "files" / "Office-3.xml"


class TestCTS(unittest.TestCase):
    def test_cheapest_package_of_measures_scenario(self) -> None:
        # Setup
        facility_etree = etree.XML("""
            <Facility xmlns="http://buildingsync.net/schemas/bedes-auc/2019">
                <Reports>
                    <Report>
                        <Scenarios>
                            <Scenario ID="Baseline"><ScenarioType><CurrentBuilding /></ScenarioType></Scenario>
                            <Scenario ID="Expensive"><ScenarioType><PackageOfMeasures><PackageFirstCost>300</PackageFirstCost></PackageOfMeasures></ScenarioType></Scenario>
                            <Scenario ID="Cheapest"><ScenarioType><PackageOfMeasures><PackageFirstCost>100</PackageFirstCost></PackageOfMeasures></ScenarioType></Scenario>
                            <Scenario ID="NoCost"><ScenarioType><PackageOfMeasures /></ScenarioType></Scenario>
                            <Scenario ID="Middle"><ScenarioType><PackageOfMeasures><PackageFirstCost>200</PackageFirstCost></PackageOfMeasures></ScenarioType></Scenario>
                        </Scenarios>
                    </Report>
                </Reports>
            </Facility>
        """)

        # Action
        scenario = FacilityAppearance(facility_etree, Path("facility.xml")).cheapest_package_of_measures_scenario

        # Assertion
        assert scenario is not None
        assert scenario.id == "Cheapest"

    def test_aggregate_facilities(self) -> None:
        # Action
        facilities = aggregate_facilities([PRIMARYSCHOOL_1_FILE_PATH, PRIMARYSCHOOL_2_FILE_PATH, OFFICE_3_FILE_PATH])

        # Assertion
        assert len(facilities["PSC MD0300ZZ"].appearances) == 2
        assert len(facilities["GSA ABC123"].appearances) == 1

    def test_cts(self) -> None:
        # Action
        with tempfile.TemporaryDirectory() as temp_dir:
            building_sync_to_cts(
                [PRIMARYSCHOOL_1_FILE_PATH, PRIMARYSCHOOL_2_FILE_PATH, OFFICE_3_FILE_PATH],
                Path(temp_dir) / "output.xlsx",
            )
            populated_cts = pd.read_excel(Path(temp_dir) / "output.xlsx")

        # Assertion
        assert len(populated_cts.index) == 5
        # top 3 row template. dtypes aren't compared since they're inferred per column on read, and the populated rows
        # can change them (e.g. a column of strings in the template becomes `object` once booleans are added)
        assert_frame_equal(populated_cts[0:3], BLANK_CTS, check_dtype=False)

    def test_cts_formatting(self) -> None:
        # Action
        with tempfile.TemporaryDirectory() as temp_dir:
            building_sync_to_cts(
                [PRIMARYSCHOOL_1_FILE_PATH, PRIMARYSCHOOL_2_FILE_PATH, OFFICE_3_FILE_PATH],
                Path(temp_dir) / "output.xlsx",
            )
            workbook = load_workbook(Path(temp_dir) / "output.xlsx")

        # Assertion
        assert workbook.sheetnames == ["Sheet1"]
        sheet = workbook["Sheet1"]
        assert sheet.auto_filter.ref == "A1:IJ1"
        assert sheet.column_dimensions["C"].width == pytest.approx(47.14785714285715)
        assert [sheet.row_dimensions[row].height for row in range(1, 6)] == [17.25, 55.5, 207.0, 40.5, None]

        # template rows keep the template's formatting
        template_cell = sheet["A3"]
        assert (template_cell.font.name, template_cell.font.sz, template_cell.font.b) == ("Arial", 10.0, True)
        assert template_cell.font.color.rgb == "FFff0000"
        assert (template_cell.fill.fill_type, template_cell.fill.fgColor.rgb) == ("solid", "FF8eb4e3")
        assert template_cell.border.left.style == "thin"

        # theme colors become black and border colors become black
        converted_cell = sheet["C4"]
        assert converted_cell.font.color.rgb == "00000000"
        assert converted_cell.border.left.color.rgb == "00000000"

        # populated rows get the default style
        data_cell = sheet["A5"]
        assert data_cell.value == "PSC MD0300ZZ"
        assert (data_cell.font.name, data_cell.font.sz, data_cell.font.b) == ("Arial", 12.0, False)
        assert (data_cell.fill.fill_type, data_cell.fill.fgColor.rgb) == ("solid", "00FFFFFF")
        assert {data_cell.border.left.style, data_cell.border.right.style, data_cell.border.top.style, data_cell.border.bottom.style} == {
            "thin",
        }
        assert (data_cell.alignment.horizontal, data_cell.alignment.vertical, data_cell.alignment.wrap_text) == ("center", "center", True)
        assert data_cell.protection.locked is False
